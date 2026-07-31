import os
import sys
import json
import logging
from docx import Document
from braille_app.doc_extractor import DocumentExtractor  # updated import

# Configure python logging to display warnings to console
logging.basicConfig(
    level=logging.WARNING,
    format="%(asctime)s - %(levelname)s - %(message)s",
    stream=sys.stdout
)

# Output directory for generated test files
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Fixtures directory for pre-built test input files
FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")

def create_mock_docx(file_path: str):
    doc = Document()

    # 1. Heading 1
    doc.add_heading("Chapter One", level=1)

    # 2. Body paragraph
    doc.add_paragraph("It was a dark and stormy night. The wind howled through the trees.")

    # 3. Bullet List Items
    doc.add_paragraph("First item in the list", style="List Bullet")
    doc.add_paragraph("Second item in the list", style="List Bullet")

    # 4. Explicit page break
    doc.add_page_break()

    # 5. Normal text on new page
    doc.add_heading("Section Two", level=2)
    doc.add_paragraph("This paragraph resides on Page 2 after the explicit page break.")

    # 6. Paragraph with custom/unknown style to trigger defensive warnings
    # Intense Quote is a standard Word style not in our mapping — intentional.
    doc.add_paragraph(
        "This is a special quote block which will trigger the unknown style warning.",
        style="Intense Quote"
    )

    doc.save(file_path)
    print(f"Created mock DOCX file: {os.path.abspath(file_path)}")

def test_extractor():
    # Generated mock file goes into tests/output/ (gitignored)
    file_path = os.path.join(OUTPUT_DIR, "mock_test_doc.docx")
    create_mock_docx(file_path)

    print("\n" + "=" * 60)
    print("RUNNING DOCUMENT EXTRACTOR VERIFICATION")
    print("=" * 60)

    extractor = DocumentExtractor()
    result = extractor.extract(file_path)

    # Print the parsed JSON
    print("\nParsed JSON Output:")
    print(json.dumps(result, indent=2, ensure_ascii=False))

    # Scan for "unknown" and "empty" style tags
    unknown_blocks = [b for p in result["pages"] for b in p["blocks"] if b["type"] == "unknown"]
    empty_blocks   = [b for p in result["pages"] for b in p["blocks"] if b["type"] == "empty"]

    print("\n" + "=" * 60)
    print("DEFENSIVE CHECK SUMMARY")
    print("=" * 60)
    if unknown_blocks:
        print(f"Detected {len(unknown_blocks)} unknown style blocks:")
        for b in unknown_blocks:
            print(f"  - Style: '{b['raw_style']}' | Text: '{b['text']}'")
    else:
        print("No unknown styles detected.")

    print(f"Detected {len(empty_blocks)} empty paragraph(s) (now preserved in output).")

    # Clean up generated file
    if os.path.exists(file_path):
        os.remove(file_path)
        print("\nCleaned up mock DOCX file.")

if __name__ == "__main__":
    test_extractor()
