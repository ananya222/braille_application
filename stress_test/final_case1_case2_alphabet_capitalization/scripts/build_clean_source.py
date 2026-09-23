"""Build the deterministic clean Case 1 plus Case 2 source master."""

from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "source" / "final_case1_case2_clean_50page.docx"
MANIFEST = ROOT / "manifests" / "source_coverage_manifest.csv"
SOURCE_RE = re.compile(r"[A-Za-z ]+")

LOWER = "abcdefghijklmnopqrstuvwxyz"
UPPER = "A B C D E F G H I J K L M N O P Q R S T U V W X Y Z"

CAP_VARIANTS = (
    ("hello", "Hello", "HELLO"),
    ("braille", "Braille", "BRAILLE"),
    ("validator", "Validator", "VALIDATOR"),
    ("english", "English", "ENGLISH"),
    ("computer", "Computer", "COMPUTER"),
    ("accessibility", "Accessibility", "ACCESSIBILITY"),
)

INTERNAL = ("iPhone", "eBay", "McDonald")


def page_rows(cycle: int) -> list[dict[str, object]]:
    variant = CAP_VARIANTS[cycle % len(CAP_VARIANTS)]
    lower, capital, upper = variant
    return [
        {
            "section": "Alphabet lowercase",
            "lines": [
                LOWER,
                LOWER + " " + LOWER,
                "a i am an as at be by do go he if in is it me my no of on or so to up us we",
                "hello braille validator ordinary education computer accessibility",
                "bookkeeper committee success address mississippi coffee green letter",
            ],
            "family": "alphabet ordinary words",
            "mode": "lowercase",
            "edge": "complete alphabet and repeated letters",
            "citation": "UEB 4.1.1",
            "treatment": "Liblouis direct uncontracted letters and words",
        },
        {
            "section": "Alphabet uppercase",
            "lines": [
                UPPER,
                "A B D Z A B D Z A B D Z",
                "A I O U B C D E F G H J K L M N P Q R S T V W X Y Z",
                "Hello Braille Validator Education",
                "NASA UEB PDF ABC",
            ],
            "family": "uppercase alphabet isolated capitals",
            "mode": "capital letters and capital words",
            "edge": "repeated isolated capitals",
            "citation": "UEB 8.1.1 8.3.1",
            "treatment": "Liblouis direct capital letter and contextual capital modes",
        },
        {
            "section": "Word lengths and repeats",
            "lines": [
                "a i am an as at be by do go he if in is it me my no of on or so to up us we",
                "cat car can cap dog dot day dry form from farm firm test text tent",
                "ordinary education computer accessibility responsibility",
                "bookkeeper committee success address mississippi assessment effective",
                "hello hello hello hello hello",
            ],
            "family": "word lengths repeated letters repeated words",
            "mode": "lowercase",
            "edge": "short long and repeated letter words",
            "citation": "UEB 4.1.1",
            "treatment": "Liblouis direct uncontracted word stream",
        },
        {
            "section": "Capitalized variants",
            "lines": [
                f"{lower} {capital} {upper}",
                f"{lower} {capital} {lower} {capital} {lower}",
                f"{capital} {capital} {capital} {capital} {capital}",
                f"{upper} {upper} {upper} {upper}",
                f"{capital} {lower} {upper} {lower} {capital}",
            ],
            "family": "same word capitalization variants",
            "mode": "lowercase capitalized all capital",
            "edge": "near neighbor capitalization",
            "citation": "UEB 8.1.1 8.3.1 8.4.1",
            "treatment": "Liblouis direct mode transitions",
        },
        {
            "section": "All capital modes",
            "lines": [
                "NASA UEB PDF ABC",
                "THIS IS ALL CAPS",
                "NASA hello hello NASA",
                "hello NASA UEB hello",
                "ABC ABC ABC abc ABC",
            ],
            "family": "all cap words capital passage",
            "mode": "capitalized word and capitalized passage",
            "edge": "mode entry continuation return and terminator",
            "citation": "UEB 8.4.1 8.4.2 8.5.1 8.5.2 8.5.3 8.6.1",
            "treatment": "Liblouis direct contextual capital modes",
        },
        {
            "section": "Internal capitals",
            "lines": [
                "iPhone eBay McDonald",
                "McDonald iPhone eBay McDonald",
                "iPhone iPhone eBay eBay McDonald McDonald",
                "iphone eBay iphone McDonald iphone",
                "McDonald NASA McDonald hello",
            ],
            "family": "internal capitals",
            "mode": "single capital indicators and capital word mode",
            "edge": "internal indicator near beginning middle and repeated",
            "citation": "UEB 8.1.1 8.3.1 8.4.1",
            "treatment": "Liblouis direct verified internal capital indicators",
        },
        {
            "section": "Similar words",
            "lines": [
                "cat car can cap",
                "form from farm firm",
                "test text tent",
                "Test test TEST Text text TEXT",
                "cat Cat CAT car Car CAR can Can CAN",
            ],
            "family": "similar word alignment",
            "mode": "lowercase capitalized all capital",
            "edge": "one letter neighbors and capitalization neighbors",
            "citation": "UEB 4.1.1 8.1.1 8.3.1 8.4.1",
            "treatment": "Liblouis direct; existing continuous alignment",
        },
        {
            "section": "Shared prefixes and suffixes",
            "lines": [
                "validate validator validation",
                "educate education educational",
                "access accessible accessibility",
                "validation Validator VALIDATION validator",
                "education Education EDUCATION educational",
            ],
            "family": "shared prefix and suffix alignment",
            "mode": "lowercase capitalized all capital",
            "edge": "long common prefixes and suffixes",
            "citation": "UEB 4.1.1 8.1.1 8.3.1 8.4.1",
            "treatment": "Liblouis direct; existing provenance mapping",
        },
        {
            "section": "Dense interaction",
            "lines": [
                "hello Hello HELLO hello Hello HELLO braille Braille BRAILLE braille Braille",
                "A A a A a A ABC ABC ABC abc ABC NASA nasa NASA UEB ueb UEB",
                "validator Validator VALIDATOR education Education EDUCATION computer Computer COMPUTER",
                "accessibility Accessibility ACCESSIBILITY iPhone eBay McDonald NASA hello Hello",
                "bookkeeper Bookkeeper BOOKKEEPER committee Committee COMMITTEE success Success SUCCESS",
            ],
            "family": "dense combined interaction",
            "mode": "all supported capitalization modes",
            "edge": "long logical stream and repeated mode changes",
            "citation": "UEB 4.1.1 8.1.1 8.3.1 8.4.1 8.5.1 8.6.1",
            "treatment": "Liblouis direct with existing alignment and provenance",
        },
        {
            "section": "Sparse boundaries",
            "lines": [
                "A",
                "a",
                "Hello",
                "hello",
                "NASA",
                "nasa",
                "iPhone",
                "McDonald",
                "ABC",
                "abc",
            ],
            "family": "sparse boundary transitions",
            "mode": "isolated capital capitalized all capital lowercase internal",
            "edge": "first and final token with sparse separation",
            "citation": "UEB 4.1.1 8.1.1 8.3.1 8.4.1 8.5.1 8.6.1",
            "treatment": "Liblouis direct block-local translation",
        },
    ]


def build() -> None:
    pages: list[dict[str, object]] = []
    for cycle in range(5):
        for family in page_rows(cycle):
            pages.append(family)
    assert len(pages) == 50
    for page in pages:
        for line in page["lines"]:  # type: ignore[index]
            if SOURCE_RE.fullmatch(str(line)) is None:
                raise ValueError(f"Unsupported source character in {line!r}")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    for page_number, page in enumerate(pages, start=1):
        if page_number > 1:
            doc.add_page_break()
        for line_number, line in enumerate(page["lines"]):  # type: ignore[index]
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(5 if line_number else 9)
            run = paragraph.add_run(str(line))
            run.font.name = "Arial"
            run.font.size = Pt(15 if line_number == 0 else 11)
            if line_number == 0:
                run.bold = True
                run.font.color.rgb = None

    DOCX.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    with MANIFEST.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=(
            "source_section", "page", "intended_test_family", "tokens_exercised",
            "capitalization_mode", "edge_case", "ueb_citation",
            "expected_liblouis_treatment", "custom_case2_handling",
        ))
        writer.writeheader()
        for page_number, page in enumerate(pages, start=1):
            writer.writerow({
                "source_section": page["section"],
                "page": page_number,
                "intended_test_family": page["family"],
                "tokens_exercised": " | ".join(page["lines"]),
                "capitalization_mode": page["mode"],
                "edge_case": page["edge"],
                "ueb_citation": page["citation"],
                "expected_liblouis_treatment": page["treatment"],
                "custom_case2_handling": "none",
            })
    print(DOCX)
    print(MANIFEST)


if __name__ == "__main__":
    build()
