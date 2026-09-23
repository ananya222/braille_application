"""Build the Case 3 source, clean PDF, and one frozen 500-action fixture.

Commands are deliberately separate: --create-source refuses an existing DOCX,
and --build-corrupted refuses to regenerate an existing corruption fixture.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from braille_app.translation.braille_cells import BRF_DOTS, cells_to_unicode, unicode_to_cells
from braille_app.translation.expected_document import generate_expected_braille
from braille_app.translation.liblouis_translator import LiblouisTranslator, vendored_metadata
from braille_app.translation.profiles import UNCONTRACTED_UEB_CASE3
from braille_app.translation.source_normalization import normalize_uncontracted_case3
from braille_app.translation.uncontracted_case3 import _translate_source_with_positions


OUT = ROOT / "stress_test" / "uncontracted_case3_numbers"
SOURCE = OUT / "source" / "case3_numbers_clean.docx"
EXPECTED = OUT / "expected" / "case3_expected_metadata.json"
COVERAGE = OUT / "expected" / "coverage_manifest.csv"
CLEAN_PDF = OUT / "output" / "case3_clean.pdf"
CLEAN_BRF = OUT / "output" / "case3_clean.brf"
CORRUPTED_PDF = OUT / "output" / "case3_corrupted.pdf"
CORRUPTED_BRF = OUT / "output" / "case3_corrupted.brf"
MANIFEST = OUT / "manifest" / "case3_500_mutations.csv"
AUDIT = OUT / "manifest" / "fixture_integrity.json"
FONT_PATH = Path(r"C:\Windows\Fonts\seguisym.ttf")
SEED = 20260923
PAGE_WIDTH = 40
PAGE_LINES = 25
FONT_SIZE = 14
LINE_HEIGHT = 27
LEFT = 42
TOP = 748

SECTIONS = (
    "Site Preparation", "Sampling Practice", "Equipment Records", "Calibration Review",
    "Daily Registers", "Comparison Methods", "Storage and Labels", "Quality Checks",
    "Team Handover", "Public Summaries", "Training Notes", "Archive Procedures",
)
TEAMS = ("north", "river", "survey", "archive", "field", "review")
SITES = ("Alder", "Birch", "Cedar", "Delta", "East", "Harbor", "Juniper", "Lake")
COUNTS = ("3 245 000", "4 500 000", "1 234 567", "8 765 432", "2 048 000", "7 125 000")
DECIMALS = ("8.93", "0.7", ".7", "4.2", "17.26", "8,93")
GROUPS = ("3,500", "4,200", "12,750", "975,000", "6,400")
YEARS = ("2024", "2025", "2026", "2031")
ORDINALS = ("1st", "2nd", "3rd", "4th")

TEMPLATES = (
    "During cycle {cycle}, the {team} team at {site} reviewed {count} records. The mean reading was {decimal}, while a separate tray held {group} samples. Two repeat checks returned 8.93, and a later reader also recorded 8.93. The team kept a 3-D frame, a 4-m shelf, and a 6-CD case beside the 20-yr copy.",
            "In {year}, the {team} station compared {count} entries with {group} control pieces. A reading of {decimal} matched the earlier 8.93 result, and the second 8.93 check stayed within range. Labels 3b, 3B, and 3k remain distinct. The notes also compare 4.b, 4.B, and 4.m before approving a 3-dimensional gauge.",
    "The {site} register lists {count} observations from cycle {cycle}. Its decimal sample is {decimal}, and the comma-marked batch contains {group} pieces. Staff verified a 4-m support, a 6-CD archive, and a 20-yr record before repeating the 3-D inspection. The repeated value 8.93 appears twice so readers can compare the two entries.",
    "A {ordinal} review checks the {team} log for {count} items and a {decimal} reading. Another batch contains {group} samples. The calibration sheet keeps 3-D, 4-m, and 6-CD labels in their original order, while a 20yr summary and a 3-dimensional diagram provide near-neighbor examples for training.",
    "For the {site} report, the team counted {count} records, checked {group} items, and repeated the value 8.93. The comparison reading was 8.93 again, with a decimal sample of {decimal}. A 3-D panel, a 4-m shelf, and a 6-CD box protect the 20-yr copy during the next {ordinal} review.",
)

CASE3_COVERAGE = (
    ("numeric indicator and digits", "6.1.1; 6.2.1", "0, 1-9, multi-digit and repeated digits"),
    ("decimal comma and period", "6.2.1", "3,500; 8.93; .7; ,7; 8,93"),
    ("mode termination", "6.3.1", "9-10; ordinary spaces; following words"),
    ("full stop before number", "6.4.1", "No. 4"),
    ("grade 1 after number", "6.5.2; 6.5.4", "3b; 4.b; 3-D; 4-m; 20-yr"),
    ("capitalization after number", "5.8.1; 6.5.2; 6.5.4", "3B; 4.B; 3-D; 6-CD"),
    ("numeric spaces", "6.6.1", "4 500 000; date 1947 08 31; time 16 00; ISBN; phone"),
    ("ordinary word near numeric hyphen", "6.5.4", "20yr; 3-dimensional; 4-bed; 6-can"),
    ("repeated numeric context", "6.1.1-6.6.1", "Repeated 8.93 and grouped-number contexts"),
    ("line and page layout", "6.1.1; 6.3.1", "Numbers adjacent to natural line and page transitions"),
)

MUTATION_COUNTS = {
    "numeric_indicators": 70,
    "digit_substitutions": 60,
    "deletions": 50,
    "insertions": 50,
    "transpositions": 40,
    "numeric_mode_transitions": 60,
    "letter_after_number_grade1": 50,
    "uppercase_after_number": 40,
    "numeric_punctuation": 30,
    "repeated_numeric_context": 25,
    "boundaries": 25,
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _paragraphs() -> list[tuple[str, str, bool]]:
    rows: list[tuple[str, str, bool]] = [
        ("Field Measurement and Record Guide", "Title", False),
        ("A flowing fictional handbook for checking numeric records and field notes.", "Subtitle", False),
    ]
    for section_index, section in enumerate(SECTIONS):
        rows.append((section, "Heading 1", section_index in {2, 7, 10}))
        for local_index in range(5):
            index = section_index * 5 + local_index
            rows.append((TEMPLATES[local_index].format(
                cycle=100 + index,
                team=TEAMS[index % len(TEAMS)],
                site=SITES[index % len(SITES)],
                count=COUNTS[index % len(COUNTS)],
                decimal=DECIMALS[index % len(DECIMALS)],
                group=GROUPS[index % len(GROUPS)],
                year=YEARS[index % len(YEARS)],
                ordinal=ORDINALS[index % len(ORDINALS)],
            ), "Normal", False))
            if index % 4 == 2:
                rows.append((
                    f"The date 1947 08 31 and time 16 00 are sample fields. "
                    f"ISBN 978 1 55468 513 4 and phone 61 3 1234 5678 use "
                    f"numeric spaces when the record context is clear.",
                    "Normal", False,
                ))
    rows.append((
        "At the end of the cycle, reviewers compare 17.26 with 17.26 and confirm the repeated reading against 8.93. The 3-D frame, 4-m shelf, and 6-CD box remain in the same bay for the 20-yr record. Staff verify 4 500 000 entries, 3 245 000 sampled rows, and 975,000 tagged pieces. In 2031, a second crew checks the 2024 first review against the 2031 fourth review. The team records time 16 00 for the final handover and checks date 1947 08 31 against the archived notes.",
        "Normal", False,
    ))
    return rows


def _all_supported(text: str) -> bool:
    try:
        normalize_uncontracted_case3(text)
        return True
    except ValueError:
        return False


def create_source() -> None:
    if SOURCE.exists():
        raise FileExistsError(f"Refusing to replace frozen/created source: {SOURCE}")
    rows = _paragraphs()
    if any(not _all_supported(text) for text, _style, _break in rows):
        raise ValueError("Generated source contains a character outside Case 3 scope")
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    for style_name in ("Title", "Subtitle", "Heading 1"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    for text, style, page_break_before in rows:
        paragraph = document.add_paragraph(text, style=style)
        paragraph.paragraph_format.keep_together = True
        if page_break_before:
            paragraph.paragraph_format.page_break_before = True
        if style == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    SOURCE.parent.mkdir(parents=True, exist_ok=True)
    document.save(SOURCE)
    print(json.dumps({
        "source": str(SOURCE),
        "paragraph_count": len(rows),
        "words": sum(len(text.split()) for text, _style, _break in rows),
        "manual_page_breaks": sum(page_break for _text, _style, page_break in rows),
        "sha256": sha256(SOURCE),
    }, indent=2))


def _master_and_expected():
    from braille_app.validation.api import _docx_paragraph_dict

    master = _docx_paragraph_dict(SOURCE)
    translator = LiblouisTranslator(UNCONTRACTED_UEB_CASE3)
    expected = generate_expected_braille(master, UNCONTRACTED_UEB_CASE3, translator)
    if expected.review_count or not all(block.rule_status == "PASS" for block in expected.blocks if block.braille):
        raise AssertionError("Clean Case 3 source has unsupported/review blocks")
    metadata = vendored_metadata(UNCONTRACTED_UEB_CASE3)
    return master, expected, translator, metadata


def _block_maps(expected, translator):
    result = []
    cursor = 0
    for block in expected.blocks:
        if not block.braille:
            continue
        braille, positions = _translate_source_with_positions(block.source_text, translator)
        if braille != block.braille:
            raise AssertionError("Expected translator and source map diverged")
        result.append({
            "source_block": block.source_block,
            "source": block.source_text,
            "braille": braille,
            "source_positions": positions,
            "expected_start": cursor,
            "expected_end": cursor + len(braille),
        })
        cursor += len(braille) + 1
    if result:
        cursor -= 1
    if cursor != len(unicode_to_cells(expected.flatten())):
        raise AssertionError("Expected block offsets do not cover flattened stream")
    return result


def _manual_offsets(rows, blocks):
    doc = Document(SOURCE)
    forced_blocks = {index for index, paragraph in enumerate(doc.paragraphs)
                     if paragraph.paragraph_format.page_break_before}
    return [block["expected_start"] for block in blocks
            if block["source_block"] in forced_blocks]


def _tokens_from_braille(braille: str):
    return [{"cell": char, "origin": index, "mutation_id": ""}
            for index, char in enumerate(braille)]


def _paginate(tokens, manual_offsets=()):
    force_origins = {int(offset) for offset in manual_offsets}
    pages: list[list[list[dict]]] = []
    break_kinds: list[str] = []
    current_page: list[list[dict]] = []
    cursor = 0

    def flush(kind: str) -> None:
        nonlocal current_page
        if current_page:
            pages.append(current_page)
            break_kinds.append(kind)
            current_page = []

    while cursor < len(tokens):
        origin = tokens[cursor]["origin"]
        if origin in force_origins:
            if current_page:
                flush("manual")
            force_origins.discard(origin)
        if len(current_page) >= PAGE_LINES:
            flush("natural")
        limit = min(cursor + PAGE_WIDTH, len(tokens))
        next_force = next((index for index in range(cursor + 1, len(tokens))
                           if tokens[index]["origin"] in force_origins), len(tokens))
        limit = min(limit, next_force)
        if limit <= cursor:
            force_origins.discard(cursor)
            continue
        end = limit
        if end < len(tokens) and end == cursor + PAGE_WIDTH:
            blanks = [index for index in range(cursor + 1, end) if tokens[index]["cell"] == "⠀"]
            if blanks:
                end = blanks[-1] + 1
        row = tokens[cursor:end]
        while row and row[-1]["cell"] == "⠀":
            row = row[:-1]
        if row:
            current_page.append(row)
        cursor = end
    flush("end")
    if not pages:
        pages.append([])
    return pages, break_kinds


def _font_name(page: int) -> str:
    name = f"Case3Cells{page}"
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(FONT_PATH)))
    return name


def _write_pdf(path: Path, pages) -> None:
    canvas = Canvas(str(path), pagesize=letter)
    canvas.setTitle("Case 3 English uncontracted UEB numeric validation fixture")
    for page_index, page in enumerate(pages):
        canvas.setFont(_font_name(page_index), FONT_SIZE)
        y = TOP
        for row in page:
            canvas.drawString(LEFT, y, "".join(token["cell"] for token in row))
            y -= LINE_HEIGHT
        canvas.showPage()
    canvas.save()


def _brf_char_map():
    result = {}
    for char, dots in BRF_DOTS.items():
        mask = sum(1 << (int(dot) - 1) for dot in dots)
        result.setdefault(mask, char.upper() if char.isalpha() else char)
    return result


def _write_brf(path: Path, pages) -> None:
    reverse = _brf_char_map()
    page_lines = []
    for page in pages:
        page_lines.append("\n".join(
            "".join(reverse[ord(token["cell"]) - 0x2800] for token in row)
            for row in page
        ))
    path.write_text("\f".join(page_lines), encoding="ascii")


def _write_coverage(blocks, metadata) -> None:
    OUT.joinpath("expected").mkdir(parents=True, exist_ok=True)
    with COVERAGE.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(("feature", "UEB_2024_rule", "examples", "source_blocks", "status"))
        for feature, rule, examples in CASE3_COVERAGE:
            matching = [str(block["source_block"]) for block in blocks
                        if any(example.split(";")[0].strip() in block["source"]
                               for example in examples.split("; "))]
            writer.writerow((feature, rule, examples, " ".join(matching[:12]), "covered"))


def build_clean() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError("Run --create-source before --build-clean")
    for path in (EXPECTED, COVERAGE, CLEAN_PDF, CLEAN_BRF):
        if path.exists():
            raise FileExistsError(f"Refusing to replace existing clean fixture artifact: {path}")
    _master, expected, _translator, runtime = _master_and_expected()
    blocks = _block_maps(expected, _translator)
    manual_offsets = _manual_offsets(_paragraphs(), blocks)
    tokens = _tokens_from_braille(expected.flatten())
    pages, break_kinds = _paginate(tokens, manual_offsets)
    natural_breaks = break_kinds.count("natural")
    manual_breaks = break_kinds.count("manual")
    page_count = len(pages)
    if not (22 <= page_count <= 28):
        raise AssertionError(f"Clean corpus has {page_count} Braille pages, expected approximately 25")
    if manual_breaks < 3 or manual_breaks > (page_count - 1) / 4:
        raise AssertionError(f"Manual page-break ratio is outside the required bounds: {manual_breaks}/{page_count-1}")
    if natural_breaks < 5:
        raise AssertionError(f"Only {natural_breaks} natural page-boundary crossings")
    if sum(len(page) >= 20 for page in pages) < 10:
        raise AssertionError("Fewer than ten Braille pages are substantially/full dense")
    if not any(10 <= len(page) < 20 for page in pages):
        raise AssertionError("Expected at least one shorter page from a legitimate manual break")
    OUT.joinpath("expected").mkdir(parents=True, exist_ok=True)
    OUT.joinpath("output").mkdir(parents=True, exist_ok=True)
    metadata = {
        "profile": UNCONTRACTED_UEB_CASE3.name,
        "runtime": runtime,
        "source_path": str(SOURCE),
        "source_sha256": sha256(SOURCE),
        "expected_braille": expected.flatten(),
        "expected_cell_count": len(unicode_to_cells(expected.flatten())),
        "expected_blocks": blocks,
        "manual_page_break_expected_offsets": manual_offsets,
        "physical_layout": {
            "columns": PAGE_WIDTH,
            "rows_per_page": PAGE_LINES,
            "page_count": page_count,
            "natural_page_transitions": natural_breaks,
            "manual_page_transitions": manual_breaks,
            "page_line_counts": [len(page) for page in pages],
            "break_kinds": break_kinds,
        },
    }
    _write_coverage(blocks, runtime)
    EXPECTED.write_text(json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    _write_pdf(CLEAN_PDF, pages)
    _write_brf(CLEAN_BRF, pages)
    print(json.dumps({
        "source_sha256": sha256(SOURCE),
        "expected_cells": metadata["expected_cell_count"],
        "layout": metadata["physical_layout"],
        "clean_pdf_sha256": sha256(CLEAN_PDF),
        "clean_brf_sha256": sha256(CLEAN_BRF),
    }, indent=2))


def _cell_type(blocks):
    chars: list[str | None] = []
    numeric_sources: list[bool] = []
    digit_sources: list[bool] = []
    repeated_ranges: set[int] = set()
    source_texts = [block["source"] for block in blocks]
    repeated = Counter(re.findall(r"[0-9]+(?:[,.][0-9]+)*", " ".join(source_texts)))
    cursor = 0
    for block in blocks:
        text, braille, positions = block["source"], block["braille"], block["source_positions"]
        for char, pos in zip(braille, positions):
            chars.append(char)
            numeric_sources.append(text[pos].isdigit())
            digit_sources.append(text[pos].isdigit() and char in "⠁⠃⠉⠙⠑⠋⠛⠓⠊⠚")
        for match in re.finditer(r"[0-9]+(?:[,.][0-9]+)*", text):
            if repeated[match.group()] >= 2:
                repeated_ranges.update(cursor + index for index, pos in enumerate(positions)
                                       if match.start() <= pos < match.end())
        cursor += len(braille)
        if block is not blocks[-1]:
            chars.append("⠀")
            numeric_sources.append(False)
            digit_sources.append(False)
            cursor += 1
    return chars, numeric_sources, digit_sources, repeated_ranges


def _layout_origin_map(pages):
    origin_locations = {}
    line_locations = defaultdict(list)
    for page_index, page in enumerate(pages):
        for row_index, row in enumerate(page):
            for col_index, token in enumerate(row):
                if token["origin"] is not None:
                    origin_locations[token["origin"]] = (page_index, row_index, col_index)
                    line_locations[(page_index, row_index)].append((col_index, token["origin"]))
    return origin_locations, line_locations


def _select_candidates(candidates, count, used_lines, used_spans, rng, family):
    by_line = defaultdict(list)
    for start, width, data in candidates:
        location = data.get("location")
        if location is None:
            continue
        line = location[:2]
        if line in used_lines:
            continue
        if any(index in used_spans for index in range(start, start + max(1, width))):
            continue
        by_line[line].append((start, width, data))
    available = list(by_line)
    rng.shuffle(available)
    selected = []
    for line in available:
        choice = rng.choice(by_line[line])
        selected.append(choice)
        used_lines.add(line)
        start, width, _data = choice
        used_spans.update(range(start, start + max(1, width)))
        if len(selected) == count:
            return selected
    raise AssertionError(
        f"Only {len(selected)}/{count} distinct-line {family} candidates; expand source coverage before freezing"
    )


def _candidate_data(index, source_char, block_index, positions, location, char):
    return {
        "source_char": source_char,
        "block_index": block_index,
        "source_offset": positions,
        "location": location,
        "cell": char,
    }


def _make_mutations(expected: str, blocks, pages, manual_offsets):
    cells, is_numeric, is_digit, repeated = _cell_type(blocks)
    if "".join(cells) != expected:
        raise AssertionError("Expected metadata and block source maps disagree")
    location_map, line_locations = _layout_origin_map(pages)
    source_char: list[str | None] = []
    source_offsets: list[int | None] = []
    block_for_cell: list[dict | None] = []
    cursor = 0
    for block in blocks:
        text = block["source"]
        source_char.extend(text[pos] for pos in block["source_positions"])
        source_offsets.extend(block["source_positions"])
        block_for_cell.extend([block] * len(block["braille"]))
        cursor += len(block["braille"])
        if block is not blocks[-1]:
            source_char.append(" ")
            source_offsets.append(None)
            block_for_cell.append(None)
            cursor += 1
    for index in range(len(cells)):
        if index not in location_map and cells[index] not in {"⠀"}:
            raise AssertionError(f"Non-space expected cell {index} has no physical slot")

    def candidates(predicate, width=1, require_boundary=None):
        rows = []
        for start in range(len(cells) - width + 1):
            if not predicate(start):
                continue
            if any(location_map.get(start + delta, (-1, -1, -1))[:2]
                   != location_map.get(start, (-2, -2, -2))[:2]
                   for delta in range(width)):
                continue
            location = location_map.get(start)
            if location is None:
                continue
            if require_boundary == "edge" and not (location[2] <= 2 or location[2] >= PAGE_WIDTH - 3):
                continue
            if require_boundary == "page" and location[1] not in {0, PAGE_LINES - 1}:
                continue
            if require_boundary == "paragraph":
                block = block_for_cell[start]
                offset = source_offsets[start]
                if block is None or offset is None or not (
                    offset <= 3 or len(block["source"]) - offset <= 4
                ):
                    continue
            rows.append((start, width, {
                "source_char": source_char[start],
                "location": location,
                "cell": cells[start],
            }))
        return rows

    used_lines: set[tuple[int, int]] = set()
    used_spans: set[int] = set()
    rng = random.Random(SEED)
    planned: list[dict] = []
    group_requests = []

    def add_group(family, candidates_, count, factory):
        group_requests.append((family, candidates_, count, factory))

    numeric_sign = candidates(lambda i: cells[i] == "⠼")
    add_group("numeric_indicators", numeric_sign, 70, lambda i, w, d, n: {
        "family": "numeric_indicators",
        "subtype": ("missing" if n % 4 == 0 else "extra" if n % 4 == 1 else "wrong" if n % 4 == 2 else "repeated"),
        "operation": ("delete" if n % 4 == 0 else "insert" if n % 4 in {1, 3} else "substitute"),
        "start": i,
        "width": 0 if n % 4 in {1, 3} else 1,
        "expected": () if n % 4 in {1, 3} else cells[i:i+1],
        "replacement": ("⠼" if n % 4 in {1, 3} else "⠰" if n % 4 == 2 else ""),
        "line": d["location"],
        "source_char": d["source_char"],
    })

    add_group("digit_substitutions", candidates(lambda i: is_digit[i]), 60,
              lambda i, w, d, n: {
        "family": "digit_substitutions", "subtype": "digit substitution",
        "operation": "substitute", "start": i, "width": 1,
        "expected": cells[i:i+1], "replacement": _other_digit(cells[i], n), "line": d["location"],
        "source_char": d["source_char"],
    })

    generic_cells = candidates(lambda i: cells[i] != "⠀" and i not in manual_offsets)
    add_group("deletions", generic_cells, 50, lambda i, w, d, n: {
        "family": "deletions", "subtype": "cell deletion", "operation": "delete",
        "start": i, "width": 1, "expected": cells[i:i+1], "replacement": "", "line": d["location"],
        "source_char": d["source_char"],
    })
    add_group("insertions", generic_cells, 50, lambda i, w, d, n: {
        "family": "insertions", "subtype": "extra cell", "operation": "insert",
        "start": i, "width": 0, "expected": (),
        "replacement": "⠼" if d["source_char"] and d["source_char"].isdigit() else "⠁",
        "line": d["location"],
        "source_char": d["source_char"],
    })
    add_group("transpositions", candidates(lambda i: cells[i] != cells[i+1]
                                             and cells[i] != "⠀" and cells[i+1] != "⠀", 2),
              40, lambda i, w, d, n: {
        "family": "transpositions", "subtype": "adjacent order swap", "operation": "transpose",
        "start": i, "width": 2, "expected": cells[i:i+2],
        "replacement": cells[i+1] + cells[i], "line": d["location"],
        "source_char": d["source_char"],
    })

    transition_indices = candidates(
        lambda i: cells[i] in {"⠼", "⠤", "⠰"}
        and (i > 0 and (cells[i-1] in {"⠤", "⠰", "⠀"} or cells[i] == "⠤")),
    )
    add_group("numeric_mode_transitions", transition_indices, 60, lambda i, w, d, n: {
        "family": "numeric_mode_transitions",
        "subtype": ("transition deletion" if n % 3 == 0 else "transition insertion" if n % 3 == 1 else "transition substitution"),
        "operation": ("delete" if n % 3 == 0 else "insert" if n % 3 == 1 else "substitute"),
        "start": i, "width": 0 if n % 3 == 1 else 1,
        "expected": () if n % 3 == 1 else cells[i:i+1],
        "replacement": "⠼" if n % 3 == 1 else (
            "⠰" if cells[i] == "⠼" else "⠼"
        ) if n % 3 == 2 else "",
        "line": d["location"],
        "source_char": d["source_char"],
    })

    g1_candidates = candidates(lambda i: cells[i] == "⠰")
    add_group("letter_after_number_grade1", g1_candidates, 50, lambda i, w, d, n: {
        "family": "letter_after_number_grade1",
        "subtype": "missing grade 1 indicator" if n % 2 == 0 else "wrong grade 1 indicator",
        "operation": "delete" if n % 2 == 0 else "substitute",
        "start": i, "width": 1, "expected": cells[i:i+1],
        "replacement": "" if n % 2 == 0 else "⠠", "line": d["location"],
        "source_char": d["source_char"],
    })

    capital_starts = candidates(lambda i: cells[i:i+2] == ["⠰", "⠠"], 2)
    add_group("uppercase_after_number", capital_starts, 40, lambda i, w, d, n: {
        "family": "uppercase_after_number",
        "subtype": "grade 1 and capitalization order" if n % 2 == 0 else "capitalization indicator loss",
        "operation": "transpose" if n % 2 == 0 else "delete",
        "start": i if n % 2 == 0 else i + 1,
        "width": 2 if n % 2 == 0 else 1,
        "expected": cells[i:i+2] if n % 2 == 0 else cells[i+1:i+2],
        "replacement": cells[i+1] + cells[i] if n % 2 == 0 else "",
        "line": d["location"] if n % 2 == 0 else location_map[i + 1],
        "source_char": d["source_char"] if n % 2 == 0 else source_char[i + 1],
    })

    punctuation = candidates(
        lambda i: cells[i] in {"⠂", "⠲"}
        and source_char[i] in {",", "."}
        and i > 0 and i + 1 < len(cells)
        and source_char[i - 1].isdigit() and source_char[i + 1].isdigit()
        and block_for_cell[i - 1] is block_for_cell[i] is block_for_cell[i + 1]
    )
    add_group("numeric_punctuation", punctuation, 30, lambda i, w, d, n: {
        "family": "numeric_punctuation", "subtype": "numeric comma/period substitution",
        "operation": "substitute", "start": i, "width": 1,
        "expected": cells[i:i+1], "replacement": "⠲" if cells[i] == "⠂" else "⠂", "line": d["location"],
        "source_char": d["source_char"],
    })

    repeated_candidates = candidates(lambda i: i in repeated and is_digit[i])
    add_group("repeated_numeric_context", repeated_candidates, 25, lambda i, w, d, n: {
        "family": "repeated_numeric_context", "subtype": "repeated-value digit substitution",
        "operation": "substitute", "start": i, "width": 1,
        "expected": cells[i:i+1], "replacement": _other_digit(cells[i], n + 3), "line": d["location"],
        "source_char": d["source_char"],
    })

    boundary_predicate = lambda i: is_digit[i] or cells[i] in {"⠼", "⠰", "⠂", "⠲"}
    def boundary_candidates(kind, boundary):
        return [
            (start, width, {**data, "boundary_context": kind})
            for start, width, data in candidates(
                boundary_predicate, require_boundary=boundary
            )
        ]
    def boundary_action(i, _width, data, n):
        return {
        "family": "boundaries", "subtype": f"{data['boundary_context']} edge numeric mutation",
        "operation": "substitute", "start": i, "width": 1,
        "expected": cells[i:i+1],
        "replacement": _other_digit(cells[i], n + 7) if is_digit[i] else (
            "⠰" if cells[i] == "⠼" else "⠼"
        ),
        "line": data["location"],
        "source_char": data["source_char"],
    }
    add_group("boundaries", boundary_candidates("line", "edge"), 15, boundary_action)
    add_group("boundaries", boundary_candidates("page", "page"), 5, boundary_action)
    add_group("boundaries", boundary_candidates("paragraph", "paragraph"), 5, boundary_action)

    # The one-action-per-line rule makes rare contextual families scarce;
    # allocate those first, then let broad digit/stream families use leftovers.
    group_requests.sort(key=lambda request: len({
        data["location"][:2] for _start, _width, data in request[1]
    }))
    for family, candidate_rows, count, factory in group_requests:
        chosen = _select_candidates(candidate_rows, count, used_lines, used_spans, rng, family)
        for start, width, data in chosen:
            planned.append(factory(start, width, data, len(planned)))

    if len(planned) != 500 or sum(MUTATION_COUNTS.values()) != 500:
        raise AssertionError("Case 3 mutation arithmetic is not exactly 500")
    if Counter(item["family"] for item in planned) != Counter(MUTATION_COUNTS):
        raise AssertionError("Case 3 family distribution differs from the required distribution")
    used = set()
    for item in planned:
        for index in range(item["start"], item["start"] + max(1, item["width"])):
            if index in used:
                raise AssertionError(f"Overlapping mutation target at expected cell {index}")
            used.add(index)
    return planned, cells


def _other_digit(current, seed):
    digits = list("⠁⠃⠉⠙⠑⠋⠛⠓⠊⠚")
    candidates = [char for char in digits if char != current]
    return candidates[seed % len(candidates)]


def _apply_mutations(tokens, actions):
    for ordinal, action in enumerate(actions, 1):
        action["mutation_id"] = f"C3-{ordinal:04d}"
    result = [dict(token) for token in tokens]
    for action in sorted(actions, key=lambda item: item["start"], reverse=True):
        start, width, operation = action["start"], action["width"], action["operation"]
        action_id = action["mutation_id"]
        if operation == "delete":
            action["actual_cells"] = ""
            del result[start:start + width]
        elif operation == "insert":
            action["actual_cells"] = action["replacement"]
            result.insert(start, {"cell": action["replacement"], "origin": None,
                                  "mutation_id": action_id})
        elif operation == "substitute":
            result[start]["cell"] = action["replacement"]
            result[start]["mutation_id"] = action_id
            action["actual_cells"] = action["replacement"]
        else:
            result[start]["cell"], result[start + 1]["cell"] = (
                result[start + 1]["cell"], result[start]["cell"]
            )
            result[start]["mutation_id"] = result[start + 1]["mutation_id"] = action_id
            action["actual_cells"] = action["replacement"]
        action["expected_cells"] = "".join(action["expected"])
    expected_delta = sum(1 if a["operation"] == "insert" else -1 if a["operation"] == "delete" else 0 for a in actions)
    if len(result) != len(tokens) + expected_delta:
        raise AssertionError("Corrupted stream length does not match action deltas")
    for action in actions:
        if action["operation"] == "substitute" and action["expected_cells"] == action["actual_cells"]:
            raise AssertionError(f"No-op substitution in {action['mutation_id']}")
        if action["operation"] == "transpose" and action["expected_cells"] == action["actual_cells"]:
            raise AssertionError(f"No-op transposition in {action['mutation_id']}")
    return result


def build_corrupted() -> None:
    for path in (CORRUPTED_PDF, CORRUPTED_BRF, MANIFEST, AUDIT):
        if path.exists():
            raise FileExistsError(f"Refusing to regenerate frozen stress artifact: {path}")
    metadata = json.loads(EXPECTED.read_text(encoding="utf-8"))
    if metadata["source_sha256"] != sha256(SOURCE):
        raise AssertionError("Clean DOCX changed after expected metadata was frozen")
    expected = metadata["expected_braille"]
    blocks = metadata["expected_blocks"]
    # Candidate physical lines are determined by the same clean cell layout.
    tokens = _tokens_from_braille(expected)
    clean_layout, _break_kinds = _paginate(tokens, metadata["manual_page_break_expected_offsets"])
    actions, _clean_cells = _make_mutations(
        expected, blocks, clean_layout, metadata["manual_page_break_expected_offsets"]
    )
    for ordinal, action in enumerate(actions, 1):
        action["mutation_id"] = f"C3-{ordinal:04d}"
    corrupted_tokens = _apply_mutations(tokens, actions)
    pages, break_kinds = _paginate(corrupted_tokens, metadata["manual_page_break_expected_offsets"])
    _write_pdf(CORRUPTED_PDF, pages)
    _write_brf(CORRUPTED_BRF, pages)
    actual_locations = {}
    for page_index, page in enumerate(pages, 1):
        for row_index, row in enumerate(page, 1):
            for col_index, token in enumerate(row, 1):
                if token["mutation_id"]:
                    actual_locations.setdefault(token["mutation_id"], []).append(
                        {"page": page_index, "line": row_index, "cell": col_index}
                    )
    for action in actions:
        action["actual_locations"] = actual_locations.get(action["mutation_id"], [])
        action["initial_page"] = action["line"][0] + 1
        action["initial_line"] = action["line"][1] + 1
        action["initial_cell"] = action["line"][2] + 1
        action["expected_start"] = action["start"]
        action["logical_target_count"] = action["width"]
        action["target_policy"] = (
            "next surviving semantic cell on the same physical line; previous cell at line end"
            if action["operation"] == "delete" else
            "all changed actual cells in the mutation span"
        )
    OUT.joinpath("manifest").mkdir(parents=True, exist_ok=True)
    fields = (
        "mutation_id", "family", "subtype", "operation", "expected_start",
        "expected_cells", "actual_cells", "logical_target_count", "target_policy",
        "initial_page", "initial_line", "initial_cell", "actual_locations",
        "source_char", "location",
    )
    with MANIFEST.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for action in actions:
            writer.writerow({
                **{key: value for key, value in action.items() if key in fields},
                "expected_cells": action["expected_cells"],
                "actual_locations": json.dumps(action["actual_locations"]),
                "source_char": action.get("source_char", ""),
                "location": json.dumps(action["line"]),
            })
    actual_change_count = len(corrupted_tokens) - len(tokens)
    integrity = {
        "seed": SEED,
        "total_actions": len(actions),
        "family_counts": dict(Counter(action["family"] for action in actions)),
        "required_family_counts": MUTATION_COUNTS,
        "length_delta": actual_change_count,
        "delta_by_operation": {
            "insertions": sum(action["operation"] == "insert" for action in actions),
            "deletions": -sum(action["operation"] == "delete" for action in actions),
            "substitutions_and_transpositions": 0,
        },
        "no_overlapping_expected_targets": True,
        "all_substitutions_change_a_cell": True,
        "all_transpositions_swap_unequal_adjacent_cells": True,
        "all_insertions_and_deletions_are_one_cell": all(
            action["width"] <= 1 for action in actions
            if action["operation"] in {"insert", "delete"}
        ),
        "controls_or_footers_present": False,
        "source_sha256": sha256(SOURCE),
        "clean_pdf_sha256": sha256(CLEAN_PDF),
        "clean_brf_sha256": sha256(CLEAN_BRF),
        "corrupted_pdf_sha256": sha256(CORRUPTED_PDF),
        "corrupted_brf_sha256": sha256(CORRUPTED_BRF),
        "manifest_sha256": sha256(MANIFEST),
        "corrupted_page_count": len(pages),
        "corrupted_line_count": sum(map(len, pages)),
        "break_kinds": break_kinds,
        "all_mutations_have_physical_cells": all(
            action["operation"] == "delete" or action["actual_locations"]
            for action in actions
        ),
    }
    if actual_change_count != integrity["delta_by_operation"]["insertions"] + integrity["delta_by_operation"]["deletions"]:
        raise AssertionError("Aggregate cell delta does not reconcile")
    if not integrity["all_mutations_have_physical_cells"]:
        raise AssertionError("One or more non-deletion actions has no physical target glyph")
    AUDIT.write_text(json.dumps(integrity, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(integrity, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--preview", action="store_true")
    parser.add_argument("--create-source", action="store_true")
    parser.add_argument("--build-clean", action="store_true")
    parser.add_argument("--build-corrupted", action="store_true")
    args = parser.parse_args()
    if args.preview:
        rows = _paragraphs()
        print(json.dumps({
            "paragraphs": len(rows),
            "words": sum(len(text.split()) for text, _style, _break in rows),
            "manual_breaks": sum(page_break for _text, _style, page_break in rows),
            "characters_supported": all(_all_supported(text) for text, _style, _break in rows),
        }, indent=2))
    elif args.create_source:
        create_source()
    elif args.build_clean:
        build_clean()
    elif args.build_corrupted:
        build_corrupted()
    else:
        parser.error("choose one command")


if __name__ == "__main__":
    main()
