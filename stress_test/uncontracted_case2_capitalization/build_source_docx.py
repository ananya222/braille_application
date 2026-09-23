"""Build the clean Case 2 source for manual Duxbury conversion."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "source" / "case2_capitalization_clean.docx"

PAGES = (
    (
        "Case Two Capitalization",
        "Lowercase Alphabet",
        "abcdefghijklmnopqrstuvwxyz",
        "Ordinary Words",
        "hello braille validator ordinary education computer",
    ),
    (
        "Uppercase Alphabet",
        "A B C D E F G H I J K L M N O P Q R S T U V W X Y Z",
        "Isolated Capitals",
        "A B D Z",
    ),
    (
        "Capitalized Words",
        "Hello Braille Validator Education",
        "Repeated Capitalized Words",
        "Hello Hello Hello",
        "hello Hello hello",
    ),
    (
        "All Capital Words",
        "NASA UEB PDF ABC",
        "All Capital And Lowercase",
        "NASA hello",
        "hello NASA",
    ),
    (
        "Internal Capital Words",
        "iPhone eBay McDonald",
        "McDonald iPhone eBay",
        "Hello NASA UEB PDF ABC hello",
    ),
    (
        "Capital Passage",
        "THIS IS ALL CAPS",
        "NASA UEB PDF ABC",
        "A B D Z",
    ),
    (
        "Source Boundaries",
        "Hello",
        "hello",
        "NASA UEB PDF ABC",
        "Hello Braille Validator Education",
    ),
    (
        "Mixed Capitalization",
        "Hello Braille Validator Education NASA UEB PDF ABC hello",
        "iPhone eBay McDonald Hello NASA hello",
        "A B D Z Hello NASA hello",
    ),
)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(14)
    for page_index, lines in enumerate(PAGES):
        if page_index:
            doc.add_page_break()
        for line_index, line in enumerate(lines):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(14 if line_index == 0 else 8)
            run = paragraph.add_run(line)
            run.font.name = "Arial"
            run.font.size = Pt(18 if line_index == 0 else 14)
            if line_index == 0:
                run.bold = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
